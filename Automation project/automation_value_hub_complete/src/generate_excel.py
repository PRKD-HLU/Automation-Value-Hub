from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.properties import CalcProperties


DARK_BLUE = "1F4E78"
MEDIUM_BLUE = "5B9BD5"
LIGHT_BLUE = "DDEBF7"
TEAL = "DDEBF7"
GRAY = "E7E6E6"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"
GREEN = "E2F0D9"
ORANGE = "FCE4D6"
LIGHT_RED = "F4CCCC"
PURPLE = "E4DFEC"
BLACK = "000000"

THIN_GRAY = Side(style="thin", color="D9E1F2")


def _month_data(kpis: pd.DataFrame, reporting_month: str) -> pd.DataFrame:
    month = pd.Timestamp(f"{reporting_month}-01")
    return kpis[pd.to_datetime(kpis["reporting_month"]) == month].copy()


def _friendly_date(value: object) -> object:
    if value is None or pd.isna(value):
        return None
    return pd.Timestamp(value).to_pydatetime()


def _write_table(
    ws,
    dataframe: pd.DataFrame,
    start_row: int,
    start_col: int,
    table_name: str,
) -> tuple[int, int]:
    for col_offset, column in enumerate(dataframe.columns, start=start_col):
        cell = ws.cell(start_row, col_offset, column)
        cell.fill = PatternFill("solid", fgColor=DARK_BLUE)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_offset, row in enumerate(dataframe.itertuples(index=False), start=start_row + 1):
        for col_offset, value in enumerate(row, start=start_col):
            cell = ws.cell(row_offset, col_offset)
            if isinstance(value, pd.Timestamp):
                cell.value = value.to_pydatetime()
                cell.number_format = "yyyy-mm-dd"
            elif pd.isna(value):
                cell.value = None
            else:
                cell.value = value
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    end_row = start_row + max(len(dataframe), 1)
    end_col = start_col + len(dataframe.columns) - 1
    if len(dataframe) > 0:
        reference = (
            f"{get_column_letter(start_col)}{start_row}:"
            f"{get_column_letter(end_col)}{end_row}"
        )
        table = Table(displayName=table_name, ref=reference)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)
    return end_row, end_col


def _set_column_widths(ws, widths: dict[str, float]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def _apply_details_formats(ws, header_row: int, last_row: int) -> None:
    headers = {
        ws.cell(header_row, column).value: column
        for column in range(1, ws.max_column + 1)
    }
    percentage_headers = {
        "First-Pass Success Rate",
        "Completion Rate",
        "Failure Rate",
        "Execution Rate",
        "Maintenance Ratio",
    }
    date_headers = {"Reporting Month"}
    decimal_headers = {
        "Gross Saved Hours",
        "Net Saved Hours",
    }

    for header, column in headers.items():
        if header in percentage_headers:
            for row in range(header_row + 1, last_row + 1):
                ws.cell(row, column).number_format = "0.0%"
        if header in date_headers:
            for row in range(header_row + 1, last_row + 1):
                ws.cell(row, column).number_format = "mmm yyyy"
        if header in decimal_headers:
            for row in range(header_row + 1, last_row + 1):
                ws.cell(row, column).number_format = "0.00"


def generate_excel_report(
    kpis: pd.DataFrame,
    issues: pd.DataFrame,
    chart_paths: dict[str, Path],
    reporting_month: str,
    ai_summary_text: str,
    output_path: Path,
) -> None:
    """Generate a formatted monthly Excel report with formulas and embedded charts."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    monthly = _month_data(kpis, reporting_month).sort_values("automation_id")

    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation = CalcProperties(
        calcMode="auto",
        fullCalcOnLoad=True,
        forceFullCalc=True,
    )

    summary = wb.create_sheet("Executive Summary")
    details = wb.create_sheet("Automation Details")
    dq_sheet = wb.create_sheet("Data Quality")
    ai_sheet = wb.create_sheet("AI Summary")
    methodology = wb.create_sheet("Methodology")

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False

    # Automation Details sheet
    details_headers = [
        "Automation ID",
        "Automation Name",
        "Process Name",
        "Technology",
        "Frequency",
        "Owner",
        "Reporting Month",
        "Expected Runs",
        "Manual Minutes",
        "Automated Minutes",
        "Machine Runtime Minutes",
        "Clean Success Runs",
        "Completed with Rework",
        "Failed Runs",
        "Retry Count",
        "Rework Minutes",
        "Maintenance Minutes",
        "Validation Status",
        "Baseline Method",
        "Baseline Confidence",
        "Comment",
        "Records Processed",
        "Actual Process Runs",
        "Completed Runs",
        "First-Pass Success Rate",
        "Completion Rate",
        "Failure Rate",
        "Execution Rate",
        "Gross Saved Minutes",
        "Gross Saved Hours",
        "Net Saved Minutes",
        "Net Saved Hours",
        "Maintenance Ratio",
        "Health Status",
        "Health Reason",
        "Data Quality Status",
    ]

    for column, header in enumerate(details_headers, start=1):
        cell = details.cell(1, column, header)
        cell.fill = PatternFill("solid", fgColor=DARK_BLUE)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GRAY)

    source_map = {
        "Automation ID": "automation_id",
        "Automation Name": "automation_name",
        "Process Name": "process_name",
        "Technology": "technology",
        "Frequency": "frequency",
        "Owner": "owner",
        "Reporting Month": "reporting_month",
        "Expected Runs": "expected_runs_used",
        "Manual Minutes": "manual_minutes_used",
        "Automated Minutes": "automated_minutes_used",
        "Machine Runtime Minutes": "machine_runtime_minutes",
        "Clean Success Runs": "clean_success_runs",
        "Completed with Rework": "completed_with_rework_runs",
        "Failed Runs": "failed_runs",
        "Retry Count": "retry_count",
        "Rework Minutes": "rework_minutes",
        "Maintenance Minutes": "maintenance_minutes",
        "Validation Status": "validation_status",
        "Baseline Method": "baseline_method_used",
        "Baseline Confidence": "baseline_confidence",
        "Comment": "comment",
        "Records Processed": "records_processed",
        "Health Status": "health_status",
        "Health Reason": "health_reason",
        "Data Quality Status": "data_quality_status",
    }

    formula_headers = {
        "Actual Process Runs",
        "Completed Runs",
        "First-Pass Success Rate",
        "Completion Rate",
        "Failure Rate",
        "Execution Rate",
        "Gross Saved Minutes",
        "Gross Saved Hours",
        "Net Saved Minutes",
        "Net Saved Hours",
        "Maintenance Ratio",
    }

    header_col = {header: idx + 1 for idx, header in enumerate(details_headers)}

    for row_index, (_, row) in enumerate(monthly.iterrows(), start=2):
        for header, source_column in source_map.items():
            value = row.get(source_column)
            cell = details.cell(row_index, header_col[header])
            if header == "Reporting Month":
                cell.value = _friendly_date(value)
                cell.number_format = "mmm yyyy"
            elif pd.isna(value):
                cell.value = None
            else:
                cell.value = value
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        col = lambda name: get_column_letter(header_col[name])
        details.cell(row_index, header_col["Actual Process Runs"], f"={col('Clean Success Runs')}{row_index}+{col('Completed with Rework')}{row_index}+{col('Failed Runs')}{row_index}")
        details.cell(row_index, header_col["Completed Runs"], f"={col('Clean Success Runs')}{row_index}+{col('Completed with Rework')}{row_index}")
        details.cell(row_index, header_col["First-Pass Success Rate"], f"=IFERROR({col('Clean Success Runs')}{row_index}/{col('Actual Process Runs')}{row_index},0)")
        details.cell(row_index, header_col["Completion Rate"], f"=IFERROR({col('Completed Runs')}{row_index}/{col('Actual Process Runs')}{row_index},0)")
        details.cell(row_index, header_col["Failure Rate"], f"=IFERROR({col('Failed Runs')}{row_index}/{col('Actual Process Runs')}{row_index},0)")
        details.cell(row_index, header_col["Execution Rate"], f"=IFERROR({col('Actual Process Runs')}{row_index}/{col('Expected Runs')}{row_index},0)")
        details.cell(row_index, header_col["Gross Saved Minutes"], f"=({col('Manual Minutes')}{row_index}-{col('Automated Minutes')}{row_index})*{col('Completed Runs')}{row_index}")
        details.cell(row_index, header_col["Gross Saved Hours"], f"={col('Gross Saved Minutes')}{row_index}/60")
        details.cell(row_index, header_col["Net Saved Minutes"], f"={col('Gross Saved Minutes')}{row_index}-{col('Rework Minutes')}{row_index}-{col('Maintenance Minutes')}{row_index}")
        details.cell(row_index, header_col["Net Saved Hours"], f"={col('Net Saved Minutes')}{row_index}/60")
        details.cell(row_index, header_col["Maintenance Ratio"], f"=IFERROR({col('Maintenance Minutes')}{row_index}/{col('Gross Saved Minutes')}{row_index},0)")

        for header in formula_headers:
            details.cell(row_index, header_col[header]).font = Font(color=BLACK)

    last_details_row = max(2, len(monthly) + 1)
    if not monthly.empty:
        details_ref = f"A1:{get_column_letter(len(details_headers))}{last_details_row}"
        table = Table(displayName="AutomationDetailsTable", ref=details_ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        details.add_table(table)

    details.freeze_panes = "A2"
    details.auto_filter.ref = f"A1:{get_column_letter(len(details_headers))}{last_details_row}"
    _set_column_widths(
        details,
        {
            "A": 14,
            "B": 30,
            "C": 28,
            "D": 20,
            "E": 12,
            "F": 24,
            "G": 14,
            "H": 13,
            "I": 14,
            "J": 16,
            "K": 17,
            "L": 15,
            "M": 18,
            "N": 12,
            "O": 12,
            "P": 14,
            "Q": 17,
            "R": 20,
            "S": 18,
            "T": 17,
            "U": 34,
            "V": 18,
            "W": 18,
            "X": 16,
            "Y": 18,
            "Z": 14,
            "AA": 15,
            "AB": 19,
            "AC": 16,
            "AD": 18,
            "AE": 17,
            "AF": 18,
            "AG": 18,
            "AH": 18,
            "AI": 40,
            "AJ": 18,
        },
    )
    _apply_details_formats(details, 1, last_details_row)

    # Data Quality sheet
    relevant_issues = issues.copy()
    if not relevant_issues.empty:
        month_string = f"{reporting_month}-01"
        relevant_issues = relevant_issues[
            relevant_issues["reporting_month"].isin(["", month_string])
        ].copy()
    dq_display = relevant_issues.rename(
        columns={
            "issue_id": "Issue ID",
            "severity": "Severity",
            "automation_id": "Automation ID",
            "reporting_month": "Reporting Month",
            "rule_code": "Rule Code",
            "field_name": "Field",
            "issue_description": "Description",
            "is_blocking": "Blocking",
            "detected_at": "Detected At",
        }
    )
    _write_table(dq_sheet, dq_display, 1, 1, "DataQualityTable")
    dq_sheet.freeze_panes = "A2"
    _set_column_widths(
        dq_sheet,
        {
            "A": 12,
            "B": 12,
            "C": 15,
            "D": 16,
            "E": 34,
            "F": 24,
            "G": 64,
            "H": 12,
            "I": 28,
        },
    )
    for row in range(2, dq_sheet.max_row + 1):
        severity = dq_sheet.cell(row, 2).value
        if severity == "Error":
            fill = LIGHT_RED
        elif severity == "Warning":
            fill = ORANGE
        else:
            fill = LIGHT_BLUE
        for col_idx in range(1, dq_sheet.max_column + 1):
            dq_sheet.cell(row, col_idx).fill = PatternFill("solid", fgColor=fill)

    # Executive Summary sheet
    summary.merge_cells("B2:M3")
    summary["B2"] = "Automation Value & Reliability Hub"
    summary["B2"].fill = PatternFill("solid", fgColor=DARK_BLUE)
    summary["B2"].font = Font(color=WHITE, bold=True, size=20)
    summary["B2"].alignment = Alignment(horizontal="center", vertical="center")

    summary.merge_cells("B4:M4")
    summary["B4"] = f"Reporting month: {pd.Timestamp(f'{reporting_month}-01').strftime('%B %Y')}"
    summary["B4"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    summary["B4"].font = Font(bold=True, size=12)
    summary["B4"].alignment = Alignment(horizontal="center")

    last_row = max(2, len(monthly) + 1)
    col = lambda name: get_column_letter(header_col[name])
    ranges = {
        name: f"'Automation Details'!{col(name)}2:{col(name)}{last_row}"
        for name in details_headers
    }

    cards = [
        ("B6:D6", "B7:D8", "Active automations", f"=COUNTA({ranges['Automation ID']})", "0"),
        ("F6:H6", "F7:H8", "Actual process runs", f"=SUM({ranges['Actual Process Runs']})", "0"),
        ("J6:L6", "J7:L8", "Net saved hours", f"=SUM({ranges['Net Saved Hours']})", "0.00"),
        ("B10:D10", "B11:D12", "First-pass success rate", f"=IF(SUM({ranges['Actual Process Runs']})=0,0,SUM({ranges['Clean Success Runs']})/SUM({ranges['Actual Process Runs']}))", "0.0%"),
        ("F10:H10", "F11:H12", "Completion rate", f"=IF(SUM({ranges['Actual Process Runs']})=0,0,SUM({ranges['Completed Runs']})/SUM({ranges['Actual Process Runs']}))", "0.0%"),
        ("J10:L10", "J11:L12", "Failed runs", f"=SUM({ranges['Failed Runs']})", "0"),
        ("B14:D14", "B15:D16", "Maintenance hours", f"=SUM({ranges['Maintenance Minutes']})/60", "0.00"),
        ("F14:H14", "F15:H16", "Action Required", f'=COUNTIF({ranges["Health Status"]},"Action Required")', "0"),
        ("J14:L14", "J15:L16", "Monitor", f'=COUNTIF({ranges["Health Status"]},"Monitor")', "0"),
    ]

    for label_range, value_range, label, formula, number_format in cards:
        summary.merge_cells(label_range)
        summary.merge_cells(value_range)
        label_cell = summary[label_range.split(":")[0]]
        value_cell = summary[value_range.split(":")[0]]
        label_cell.value = label
        label_cell.fill = PatternFill("solid", fgColor=MEDIUM_BLUE)
        label_cell.font = Font(color=WHITE, bold=True)
        label_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.value = formula
        value_cell.fill = PatternFill("solid", fgColor=TEAL)
        value_cell.font = Font(color=BLACK, bold=True, size=18)
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.number_format = number_format

    summary.merge_cells("B19:M19")
    summary["B19"] = "Key visuals"
    summary["B19"].fill = PatternFill("solid", fgColor=DARK_BLUE)
    summary["B19"].font = Font(color=WHITE, bold=True)

    if chart_paths.get("net_saved_hours") and chart_paths["net_saved_hours"].exists():
        image = Image(str(chart_paths["net_saved_hours"]))
        image.width = 610
        image.height = 335
        summary.add_image(image, "B21")
    if chart_paths.get("health_status") and chart_paths["health_status"].exists():
        image = Image(str(chart_paths["health_status"]))
        image.width = 450
        image.height = 300
        summary.add_image(image, "H21")

    summary.merge_cells("B43:M46")
    summary["B43"] = (
        "Interpretation: Gross savings represent manual active time avoided for completed runs. "
        "Net savings additionally deduct rework and maintenance. Machine runtime is stored separately "
        "and is not deducted when the process is unattended or semi-attended."
    )
    summary["B43"].fill = PatternFill("solid", fgColor=LIGHT_GRAY)
    summary["B43"].alignment = Alignment(wrap_text=True, vertical="top")
    summary["B43"].font = Font(italic=True, color="666666")
    _set_column_widths(summary, {get_column_letter(i): 14 for i in range(1, 14)})
    summary.row_dimensions[2].height = 28
    summary.row_dimensions[3].height = 28

    # AI Summary sheet
    ai_sheet.merge_cells("A1:H2")
    ai_sheet["A1"] = "Monthly Management Summary"
    ai_sheet["A1"].fill = PatternFill("solid", fgColor=DARK_BLUE)
    ai_sheet["A1"].font = Font(color=WHITE, bold=True, size=18)
    ai_sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ai_sheet.merge_cells("A4:H24")
    ai_sheet["A4"] = ai_summary_text
    ai_sheet["A4"].alignment = Alignment(wrap_text=True, vertical="top")
    ai_sheet["A4"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    ai_sheet["A4"].font = Font(size=11)
    ai_sheet["A4"].comment = Comment(
        "This text is generated from validated KPI data. An approved AI tool may improve the wording but must not recalculate the figures.",
        "Automation Value Hub",
    )
    _set_column_widths(ai_sheet, {get_column_letter(i): 16 for i in range(1, 9)})

    # Methodology sheet
    methodology_headers = ["Metric / Concept", "Definition"]
    for col_idx, header in enumerate(methodology_headers, start=1):
        cell = methodology.cell(1, col_idx, header)
        cell.fill = PatternFill("solid", fgColor=DARK_BLUE)
        cell.font = Font(color=WHITE, bold=True)
    definitions = [
        ("Actual Process Runs", "Clean Success Runs + Completed with Rework + Failed Runs."),
        ("Completed Runs", "Clean Success Runs + Completed with Rework."),
        ("First-Pass Success Rate", "Clean Success Runs / Actual Process Runs."),
        ("Completion Rate", "Completed Runs / Actual Process Runs."),
        ("Failure Rate", "Failed Runs / Actual Process Runs."),
        ("Execution Rate", "Actual Process Runs / Expected Runs per Month."),
        ("Gross Saved Minutes", "(Manual Active Minutes - Automated Active Minutes) x Completed Runs."),
        ("Net Saved Minutes", "Gross Saved Minutes - Rework Minutes - Maintenance Minutes."),
        ("Maintenance Ratio", "Maintenance Minutes / Gross Saved Minutes."),
        ("Machine Runtime", "Technical execution time. It is not deducted from user time for unattended or semi-attended processes."),
        ("Healthy", "No failure, rework, retry or material maintenance issue was detected."),
        ("Monitor", "The process completed, but rework, retry, execution gap or elevated maintenance was recorded."),
        ("Action Required", "A failed run, failed validation, non-positive net saving or high maintenance ratio was detected."),
        ("Insufficient Data", "A required baseline or monthly process result is missing."),
    ]
    for row_idx, (metric, definition) in enumerate(definitions, start=2):
        methodology.cell(row_idx, 1, metric)
        methodology.cell(row_idx, 2, definition)
        methodology.cell(row_idx, 2).alignment = Alignment(wrap_text=True, vertical="top")
    _set_column_widths(methodology, {"A": 30, "B": 100})
    methodology.freeze_panes = "A2"

    wb.save(output_path)
